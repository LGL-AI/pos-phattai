"""Append bug rows (and rebuild the summary sheet) to docs/LotusPOS_BugList.xlsx keeping the sheet's formatting.

python3 scripts/add-bug.py docs/LotusPOS_BugList.xlsx '<json list of 15-field rows>'
"""
import json, sys, copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

path, rows = sys.argv[1], json.loads(sys.argv[2])
wb = load_workbook(path)
ws = wb['Bug List']
table = ws.tables['Bugs']
last = ws.max_row
sev_fill = {'Cao': 'F6D9D6', 'Trung': 'FBECD1', 'Thấp': 'E4ECEA'}
for row in rows:
    assert len(row) == 15, row
    last += 1
    for col, value in enumerate(row, 1):
        cell = ws.cell(row=last, column=col, value=value)
        src = ws.cell(row=2, column=col)
        cell.font = copy.copy(src.font)
        cell.alignment = copy.copy(src.alignment)
        cell.border = copy.copy(src.border)
    ws.cell(row=last, column=7).fill = PatternFill('solid', fgColor=sev_fill[row[6]])
table.ref = f"A1:O{last}"

# Rebuild the summary so new groups get a row (values are live COUNTIFS formulas).
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
stats = wb['Thống kê']
head_font, head_fill = copy.copy(stats['A1'].font), copy.copy(stats['A1'].fill)
body = Font(name='Arial')
stats.delete_rows(2, stats.max_row)
groups = sorted({ws.cell(row=r, column=8).value for r in range(2, last + 1) if ws.cell(row=r, column=8).value})
for r, g in enumerate(groups, 2):
    stats.cell(row=r, column=1, value=g).font = body
    for ci, sev in enumerate(['Cao', 'Trung', 'Thấp'], 2):
        stats.cell(row=r, column=ci, value=f"=COUNTIFS('Bug List'!$H:$H,$A{r},'Bug List'!$G:$G,\"{sev}\")").font = body
    stats.cell(row=r, column=5, value=f'=SUM(B{r}:D{r})').font = Font(name='Arial', bold=True)
tr = len(groups) + 2
stats.cell(row=tr, column=1, value='Tổng').font = Font(name='Arial', bold=True)
for ci in range(2, 6):
    col = get_column_letter(ci)
    stats.cell(row=tr, column=ci, value=f'=SUM({col}2:{col}{tr - 1})').font = Font(name='Arial', bold=True)
sr = tr + 2
for ci, h in enumerate(['Trạng thái', 'Số lỗi'], 1):
    stats.cell(row=sr, column=ci, value=h).font = Font(name='Arial', bold=True)
for i, st in enumerate(['Mới', 'Đang làm', 'Đã sửa', 'Không sửa'], 1):
    stats.cell(row=sr + i, column=1, value=st).font = body
    stats.cell(row=sr + i, column=2, value=f"=COUNTIF('Bug List'!$O:$O,A{sr + i})").font = body
stats.cell(row=sr + 6, column=1, value='Bảng này do scripts/add-bug.py dựng lại từ cột H mỗi lần thêm lỗi.').font = Font(name='Arial', italic=True, color='5D6E6C')
wb.save(path)
print('rows now', last - 1)
